"""D题问题一：单点往返安全载荷与逐箱组批。

运行：python d_q1.py [输出目录] [D题数据目录]
依赖：openpyxl、Pillow（用于读取附件 XLSX 和 GeoTIFF）。
"""
from __future__ import annotations

from collections import defaultdict
from functools import lru_cache
from math import asin, cos, floor, radians, sin, sqrt
from pathlib import Path
import json
import sys

from openpyxl import load_workbook
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DATA = (Path(sys.argv[2]) if len(sys.argv)>2 else
        Path(r"E:\数学建模\第二十三届中国研究生数学建模竞赛 - 中文题目\中文题目\D题\数据"))
BASE = DATA / "无人机应急物资运输基础数据"
G = 9.80665
EPS = 1e-9


def read_nodes():
    ws = load_workbook(BASE / "调度中心与服务区.xlsx", read_only=True, data_only=True).active
    nodes = {}
    for row in ws.values:
        if isinstance(row[0], str) and (row[0] == "O01" or row[0].startswith("S0")):
            nodes[row[0]] = dict(lon=float(row[2]), lat=float(row[3]), ground=float(row[4]))
    return nodes


def read_aircraft():
    ws = load_workbook(BASE / "运输无人机数据.xlsx", read_only=True, data_only=True).active
    aircraft = {}
    for row in ws.values:
        if row[0] not in ("A", "B", "C") or not isinstance(row[3], (int, float)):
            continue
        aircraft[row[0]] = dict(
            id=row[0], mass0=float(row[2]), maxload=float(row[3]), maxvolume=float(row[4]),
            cruise_speed=float(row[5]), range_empty=float(row[6]), range_full=float(row[7]),
            battery=float(row[8]), reserve=float(row[9]) / 100,
            prepare=float(row[10]), load_each=float(row[11]),
            handoff_base=float(row[12]), handoff_each=float(row[13]),
            climb_speed=float(row[14]), descend_speed=float(row[15]),
            climb_efficiency=float(row[16]),
        )
    return aircraft


def read_boxes():
    ws = load_workbook(BASE / "物资需求与配送时限.xlsx", read_only=True, data_only=True)["逐箱货箱清单"]
    boxes = defaultdict(list)
    for row in list(ws.values)[1:]:
        if not row[0]:
            continue
        boxes[row[1]].append(dict(id=row[0], mass=float(row[3]), volume=float(row[4]),
                                  priority=int(row[8])))
    return dict(boxes)


def haversine(lon1, lat1, lon2, lat2):
    # WGS84 mean Earth radius; over this small region, distance error is minor.
    r = 6371008.8
    a = sin(radians(lat2-lat1)/2)**2 + cos(radians(lat1))*cos(radians(lat2))*sin(radians(lon2-lon1)/2)**2
    return 2*r*asin(min(1.0,sqrt(a)))


class Dem:
    def __init__(self, path):
        self.im = Image.open(path)
        assert self.im.mode == "F"
        self.width, self.height = self.im.size
        self.dx, self.dy, _ = self.im.tag_v2[33550]
        _, _, _, self.lon0, self.lat0, _ = self.im.tag_v2[33922]
        self.pixels = self.im.load()

    def xy(self, node):
        x = (node["lon"] - self.lon0) / self.dx
        y = (self.lat0 - node["lat"]) / self.dy
        if not (0 <= x < self.width and 0 <= y < self.height):
            raise ValueError(f"Node outside DEM: {node}")
        return x, y

    def max_on_segment(self, start, end):
        """Visit each DEM cell intersected by a horizontal straight route."""
        x0,y0=self.xy(start); x1,y1=self.xy(end)
        ix,iy=floor(x0),floor(y0)
        ex,ey=floor(x1),floor(y1)
        vx,vy=x1-x0,y1-y0
        sx=1 if vx>0 else -1 if vx<0 else 0
        sy=1 if vy>0 else -1 if vy<0 else 0
        tx=((ix+1-x0)/vx if sx>0 else (ix-x0)/vx) if sx else float("inf")
        ty=((iy+1-y0)/vy if sy>0 else (iy-y0)/vy) if sy else float("inf")
        dtx=abs(1/vx) if sx else float("inf")
        dty=abs(1/vy) if sy else float("inf")
        visited=set()
        def add(a,b):
            if not (0<=a<self.width and 0<=b<self.height):
                raise ValueError("Route leaves DEM")
            visited.add((a,b))
        add(ix,iy)
        while (ix,iy)!=(ex,ey):
            if tx < ty-EPS:
                ix+=sx; tx+=dtx; add(ix,iy)
            elif ty < tx-EPS:
                iy+=sy; ty+=dty; add(ix,iy)
            else:
                add(ix+sx,iy); add(ix,iy+sy)
                ix+=sx; iy+=sy; tx+=dtx; ty+=dty; add(ix,iy)
        return max(float(self.pixels[x,y]) for x,y in visited), len(visited)


def segment_energy(ac, distance, climb, payload):
    if payload < -EPS or payload > ac["maxload"]+EPS:
        return float("inf")
    qratio=max(0.0,payload/ac["maxload"])
    equivalent_range=ac["range_empty"]-(ac["range_empty"]-ac["range_full"])*qratio**1.5
    # Horizontal energy is battery usable energy × distance / payload-dependent range.
    horizontal=ac["battery"]*distance/equivalent_range
    # Positive elevation gain requires additional gravitational energy.
    upward=(ac["mass0"]+payload)*G*max(0.0,climb)/(3.6e6*ac["climb_efficiency"])
    return horizontal+upward


def route_metrics(node_id, nodes, dem):
    origin,site=nodes["O01"],nodes[node_id]
    high,ncells=dem.max_on_segment(origin,site)
    cruise=high+50
    altitude0=origin["ground"]
    altitude1=site["ground"]+30
    if cruise+EPS<max(altitude0,altitude1):
        raise ValueError(f"Cruise altitude below operation altitude: {node_id}")
    d=haversine(origin["lon"],origin["lat"],site["lon"],site["lat"])
    return dict(service=node_id,distance=d,terrain_max=high,cruise_altitude=cruise,
                dem_cells=ncells,out_climb=cruise-altitude0,out_descend=cruise-altitude1,
                return_climb=cruise-altitude1,return_descend=cruise-altitude0)


def trip(ac, route, mass, count):
    eout=segment_energy(ac,route["distance"],route["out_climb"],mass)
    eback=segment_energy(ac,route["distance"],route["return_climb"],0.0)
    energy=eout+eback
    fly=(route["out_climb"]+route["return_climb"])/ac["climb_speed"]
    fly+=(route["out_descend"]+route["return_descend"])/ac["descend_speed"]
    fly+=2*route["distance"]/ac["cruise_speed"]
    work=ac["prepare"]+count*ac["load_each"]+fly+ac["handoff_base"]+count*ac["handoff_each"]
    return energy,work


def safe_payload(ac,route):
    capacity=(1-ac["reserve"])*ac["battery"]
    if trip(ac,route,0,0)[0]>capacity+EPS:
        return None
    if trip(ac,route,ac["maxload"],0)[0]<=capacity+EPS:
        return ac["maxload"]
    lo,hi=0.0,ac["maxload"]
    for _ in range(55):
        mid=(lo+hi)/2
        if trip(ac,route,mid,0)[0]<=capacity:
            lo=mid
        else:
            hi=mid
    return lo


def optimize_site(site, boxes, aircraft, route, priority, allowed_models=None):
    n=len(boxes)
    full=(1<<n)-1
    mass=[0.0]*(full+1); vol=[0.0]*(full+1); count=[0]*(full+1)
    for mask in range(1,full+1):
        bit=mask&-mask; k=bit.bit_length()-1; prev=mask^bit
        mass[mask]=mass[prev]+boxes[k]["mass"]
        vol[mask]=vol[prev]+boxes[k]["volume"]
        count[mask]=count[prev]+1
    candidates=[None]*(full+1)
    for mask in range(1,full+1):
        opts=[]
        for a in aircraft.values():
            if allowed_models is not None and a["id"] not in allowed_models:
                continue
            if mass[mask]>a["maxload"]+EPS or vol[mask]>a["maxvolume"]+EPS:
                continue
            energy,work=trip(a,route,mass[mask],count[mask])
            if energy <= (1-a["reserve"])*a["battery"]+EPS:
                opts.append((a["id"],energy,work))
        candidates[mask]=opts

    # Pivot on the least significant uncovered box so each unordered partition is examined once.
    @lru_cache(None)
    def solve(mask):
        if mask==0:
            return (0,0.0,0.0,())
        anchor=mask&-mask
        sub=mask
        best=None
        while sub:
            if sub&anchor and candidates[sub]:
                tail=solve(mask^sub)
                if tail is not None:
                    for model,e,t in candidates[sub]:
                        cand=(tail[0]+1,tail[1]+e,tail[2]+t,((sub,model),)+tail[3])
                        key=lambda v: tuple(v[i] for i in priority)
                        if best is None or key(cand)<key(best):
                            best=cand
            sub=(sub-1)&mask
        return best

    solution=solve(full)
    if solution is None:
        raise ValueError(f"No feasible cover for {site}")
    rows=[]
    for mask,model in solution[3]:
        ac=aircraft[model]
        e,t=trip(ac,route,mass[mask],count[mask])
        rows.append(dict(service=site,model=model,boxes=[boxes[i]["id"] for i in range(n) if mask&(1<<i)],
                         mass=mass[mask],volume=vol[mask],energy=e,work_time=t,
                         return_soc=100*(1-e/ac["battery"])))
    rows.sort(key=lambda r:(-sum(next(b["priority"] for b in boxes if b["id"]==bid) for bid in r["boxes"]),r["model"]))
    return rows


def main(outdir):
    nodes=read_nodes(); aircraft=read_aircraft(); boxes=read_boxes()
    dem=Dem(next(DATA.rglob("*.tif")))
    routes={s:route_metrics(s,nodes,dem) for s in sorted(boxes)}
    safety=[]
    for s,r in routes.items():
        for a in aircraft.values():
            m=safe_payload(a,r)
            safety.append(dict(service=s,model=a["id"],max_safe_payload_kg=m,
                               roundtrip_empty_kwh=trip(a,r,0,0)[0],
                               energy_limit_kwh=(1-a["reserve"])*a["battery"]))

    strategies={
        "sorties_then_energy_then_time":(0,1,2),
        "energy_then_sorties_then_time":(1,0,2),
        "time_then_sorties_then_energy":(2,0,1),
    }
    results={}
    for name,priority in strategies.items():
        all_rows=[]
        for s in sorted(boxes):
            all_rows+=optimize_site(s,boxes[s],aircraft,routes[s],priority)
        ids=[bid for row in all_rows for bid in row["boxes"]]
        expected={b["id"] for bs in boxes.values() for b in bs}
        assert len(ids)==len(expected)==80 and set(ids)==expected
        for row in all_rows:
            a=aircraft[row["model"]]
            assert row["mass"]<=a["maxload"]+EPS
            assert row["volume"]<=a["maxvolume"]+EPS
            assert row["energy"]<=(1-a["reserve"])*a["battery"]+EPS
        results[name]=dict(summary=dict(sorties=len(all_rows),energy_kwh=sum(r["energy"] for r in all_rows),
                                        work_time_s=sum(r["work_time"] for r in all_rows),
                                        delivered_boxes=len(ids)),trips=all_rows)

    fixed_model_results={}
    for model in aircraft:
        all_rows=[]
        for s in sorted(boxes):
            all_rows+=optimize_site(s,boxes[s],aircraft,routes[s],(0,1,2),{model})
        ids=[bid for row in all_rows for bid in row["boxes"]]
        expected={b["id"] for bs in boxes.values() for b in bs}
        assert len(ids)==len(expected)==80 and set(ids)==expected
        for row in all_rows:
            a=aircraft[row["model"]]
            assert row["model"]==model
            assert row["mass"]<=a["maxload"]+EPS
            assert row["volume"]<=a["maxvolume"]+EPS
            assert row["energy"]<=(1-a["reserve"])*a["battery"]+EPS
        fixed_model_results[model]=dict(
            summary=dict(sorties=len(all_rows),energy_kwh=sum(r["energy"] for r in all_rows),
                         work_time_s=sum(r["work_time"] for r in all_rows),delivered_boxes=len(ids)),
            trips=all_rows)

    sensitivity=[]
    for reserve in (0.10,0.20,0.30):
        adjusted={key:dict(a,reserve=reserve) for key,a in aircraft.items()}
        if reserve==0.20:
            trips=results["sorties_then_energy_then_time"]["trips"]
        else:
            trips=[]
            for s in sorted(boxes):
                trips+=optimize_site(s,boxes[s],adjusted,routes[s],(0,1,2))
        limits=[dict(service=s,model=a["id"],max_safe_payload_kg=safe_payload(adjusted[a["id"]],r))
                for s,r in routes.items() for a in aircraft.values()]
        sensitivity.append(dict(reserve_fraction=reserve,sorties=len(trips),
                                energy_kwh=sum(x["energy"] for x in trips),
                                work_time_s=sum(x["work_time"] for x in trips),
                                affected_max_loads=sum(
                                    1 for x in limits if x["max_safe_payload_kg"] is None or
                                    x["max_safe_payload_kg"] < adjusted[x["model"]]["maxload"]-1e-6),
                                limits=limits))

    outdir.mkdir(parents=True,exist_ok=True)
    payload=dict(input_counts=dict(nodes=len(nodes),services=len(boxes),boxes=sum(map(len,boxes.values())),
                                   models=len(aircraft)),assumptions=dict(
        distance="WGS84 haversine",terrain="All DEM pixels touched by the horizontal segment",
        horizontal_energy="Euse * distance / L(q)",
        climb_energy="(empty_mass_including_battery + payload) * g * positive_climb / (3.6e6 * climb_efficiency)",
        priority="Lexicographic; Q1 has no entity-aircraft or shared-battery scheduling"),
        routes=routes,safety=safety,strategies=results,fixed_model_results=fixed_model_results,
        sensitivity=sensitivity)
    (outdir/"d_q1_results.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    print("INPUT",payload["input_counts"])
    print("SAFETY MAX PAYLOAD KG")
    for s in sorted(boxes):
        print(s," ".join(f'{a}: {next(x["max_safe_payload_kg"] for x in safety if x["service"]==s and x["model"]==a):.2f}' if next(x["max_safe_payload_kg"] for x in safety if x["service"]==s and x["model"]==a) is not None else f'{a}: NONE' for a in aircraft))
    for name,v in results.items():
        print(name,v["summary"])
    for model,v in fixed_model_results.items():
        print("ALL_MODEL",model,v["summary"])
    for x in sensitivity:
        print("SENSITIVITY",{k:v for k,v in x.items() if k!="limits"})
    print("OUTPUT",outdir/"d_q1_results.json")


if __name__=="__main__":
    main(Path(sys.argv[1]) if len(sys.argv)>1 else Path(__file__).parent)
