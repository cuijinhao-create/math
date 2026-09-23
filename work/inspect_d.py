from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from openpyxl import load_workbook
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

base = Path(r"E:\数学建模\第二十三届中国研究生数学建模竞赛 - 中文题目\中文题目")
d = base / "D题"
with ZipFile(base / "D题.zip") as z:
    names = [n for n in z.namelist() if not n.endswith("/")]
    code = [n for n in names if Path(n).suffix.lower() in {".py", ".ipynb", ".m", ".r", ".jl", ".cpp", ".c", ".java", ".js", ".ts"}]
    print("ZIP files", len(names), "types", dict(Counter(Path(n).suffix.lower() for n in names)))
    print("ZIP CODE", code)
for p in d.rglob("*.xlsx"):
    if p.name.startswith("~$"):
        continue
    print("\nWORKBOOK", p.relative_to(d))
    wb = load_workbook(p, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print("SHEET", ws.title, ws.max_row, ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,4), values_only=True):
            print(" ", [str(v)[:65] if v is not None else None for v in row[:16]])
for p in d.rglob("*.csv"):
    print("CSV", p.relative_to(d), "bytes", p.stat().st_size)
